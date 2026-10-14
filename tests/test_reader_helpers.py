"""Public regression tests; no private manuscripts or renderer binaries required."""
from __future__ import annotations

from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / "skills" / "story-world-architect"
SCRIPTS = SKILL / "scripts"
sys.path.insert(0, str(SCRIPTS))
from build_reader_docx import (DEFAULT_TEMPLATE, ROLE_STYLES, SourceError, build,
                               paragraph_xml_text, parse_markdown)
from extract_reader_template import sanitize_template_metadata
from render_pdf_pages import render
from verify_reader_docx import body_stream, style_fonts, verify
import fitz


class ReaderHelpersTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="world-skill-test-")
        self.folder = Path(self.temporary.name)
        self.source = self.folder / "source.md"

    def tearDown(self):
        self.temporary.cleanup()

    def source_text(self, body):
        self.source.write_text("# 世界测试\n\n" + body + "\n", encoding="utf-8")
        return parse_markdown(self.source)

    def test_heading_conventions_are_equivalent(self):
        direct = self.source_text("# 一\n\n## 二\n\n### 三\n\n正文。").semantic_stream()
        nested = self.source_text("## 一\n\n### 二\n\n#### 三\n\n正文。").semantic_stream()
        self.assertEqual(direct, nested)

    def test_unsupported_content_fails_explicitly(self):
        cases = ("![图](x.png)", "> 引文", "```text\n代码\n```", "### 跳级",
                 "- 项目\n  - 嵌套", "{{BODY}}", "3. 项目\n5. 错误编号")
        for body in cases:
            with self.subTest(body=body), self.assertRaises(SourceError):
                self.source_text("# 章节\n\n" + body)

    def test_body_source_and_numbering_preserved(self):
        parsed = self.source_text("# 章节\n\n段落中的**强调**。\n\n## 小节\n\n3. 第一项\n4. 第二项\n\n| 甲 | 乙 |\n| --- | --- |\n| 内容 | 另一内容 |")
        out = self.folder / "book.docx"
        build(self.source, out)
        document = Document(out)
        self.assertEqual(parsed.semantic_stream(), body_stream(document))
        self.assertEqual([paragraph_xml_text(p) for p in document.paragraphs if p.style.name == "Contents Title"], ["目录"])
        starts = list(document.element.iter(qn("w:bookmarkStart")))
        self.assertEqual([e.get(qn("w:name")) for e in starts], ["reader_heading_0001", "reader_heading_0002"])

    def test_all_four_font_slots_match_template(self):
        self.source_text("# 章节\n\n正文。")
        out = self.folder / "book.docx"
        build(self.source, out)
        actual, template = Document(out), Document(DEFAULT_TEMPLATE)
        for role in ROLE_STYLES:
            with self.subTest(role=role):
                fonts = style_fonts(actual, actual.styles[role])
                self.assertTrue(all(fonts.values()))
                self.assertEqual(fonts, style_fonts(template, template.styles[role]))

    def test_output_cannot_overwrite_source_or_template(self):
        self.source_text("# 章节\n\n正文。")
        for out in (self.source, DEFAULT_TEMPLATE):
            with self.subTest(out=out), self.assertRaises(ValueError):
                build(self.source, out)

    def test_initializer_defaults_reader_and_does_not_overwrite(self):
        project = self.folder / "project"
        command = [sys.executable, "-B", "-X", "utf8", str(SCRIPTS / "init_world_project.py"),
                   "--root", str(project), "--title", "测试", "--mode", "compact"]
        subprocess.run(command, check=True, capture_output=True)
        before = {p.relative_to(project): p.read_bytes() for p in project.rglob("*.md")}
        self.assertEqual(len(before), 10)
        subprocess.run(command + ["--audience", "author"], check=True, capture_output=True)
        self.assertEqual(before, {p.relative_to(project): p.read_bytes() for p in project.rglob("*.md")})
        self.assertIn("reader", (project / "设定" / "世界观" / "世界观索引.md").read_text(encoding="utf-8"))

    def test_template_has_no_old_authoring_metadata(self):
        with ZipFile(DEFAULT_TEMPLATE) as archive:
            core = etree.fromstring(archive.read("docProps/core.xml"))
            self.assertFalse(any((e.text or "").strip() for e in core.iter()))
            app = etree.fromstring(archive.read("docProps/app.xml"))
            self.assertEqual([etree.QName(e).localname for e in app], ["Application"])
            self.assertNotIn("WPS", archive.read("docProps/app.xml").decode())
            self.assertFalse(any("customXml/" in name or name == "docProps/custom.xml" for name in archive.namelist()))

    def test_metadata_scrub_preserves_styles_and_text(self):
        source_bytes = DEFAULT_TEMPLATE.read_bytes()
        copied = self.folder / "template.docx"
        copied.write_bytes(source_bytes)
        before = Document(copied)
        sanitize_template_metadata(copied)
        after = Document(copied)
        self.assertEqual([p.text for p in before.paragraphs], [p.text for p in after.paragraphs])
        for role in ROLE_STYLES:
            self.assertEqual(before.styles[role].element.xml, after.styles[role].element.xml)
        self.assertEqual(DEFAULT_TEMPLATE.read_bytes(), source_bytes)

    def test_actual_pdf_page_mapping_and_corruption_detection(self):
        self.source_text("# Test\n\nParagraph.")
        pdf = self.folder / "pages.pdf"
        with fitz.open() as document:
            document.new_page()
            page = document.new_page()
            page.insert_text((60, 60), "Test")
            document.set_toc([[1, "Test", 2]])
            document.save(pdf)
        out = self.folder / "book.docx"
        build(self.source, out, pages_from_pdf=pdf)
        report = verify(self.source, out, pdf)
        self.assertTrue(report["ok"], report["errors"])
        damaged = Document(out)
        next(p for p in damaged.paragraphs if p.style.name == "Reading Body").text = "Changed."
        damaged.save(out)
        self.assertFalse(verify(self.source, out, pdf)["ok"])

    def test_pdf_renderer_refuses_stale_pages(self):
        pdf = self.folder / "pages.pdf"
        with fitz.open() as document:
            document.new_page()
            document.save(pdf)
        out = self.folder / "pages"
        self.assertEqual(render(pdf, out, 72), 1)
        self.assertTrue((out / "page-1.png").is_file())
        with self.assertRaises(ValueError):
            render(pdf, out, 72)


if __name__ == "__main__":
    unittest.main()
