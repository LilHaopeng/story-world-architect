"""Behavioral regression fixtures for four-level reader books and sparse endings."""
from __future__ import annotations

from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
import fitz

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skills" / "story-world-architect" / "scripts"
sys.path.insert(0, str(SCRIPTS))
from build_reader_docx import DEFAULT_TEMPLATE, SourceError, build, parse_markdown, paragraph_xml_text
from extract_reader_template import distill
from verify_reader_docx import body_stream, geometry, verify


class FourLevelReaderTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="world-h4-test-")
        self.folder = Path(self.temporary.name)
        self.source = self.folder / "source.md"

    def tearDown(self):
        self.temporary.cleanup()

    def write_source(self, text):
        self.source.write_text("# Book\n\n" + text + "\n", encoding="utf-8")
        return parse_markdown(self.source)

    def pdf(self, headings, pages=None):
        out = self.folder / "pages.pdf"
        pages = pages or [3] * len(headings)
        with fitz.open() as doc:
            for _ in range(max(pages)):
                doc.new_page()
            doc.set_toc([[level, title, page] for (level, title), page in zip(headings, pages)])
            doc.save(out)
        return out

    def test_four_levels_normalize_equivalently(self):
        direct = self.write_source("# One\n\n## Two\n\n### Three\n\n#### Four\n\nBody.")
        nested = self.write_source("## One\n\n### Two\n\n#### Three\n\n##### Four\n\nBody.")
        self.assertEqual(direct.headings, [(1, "One"), (2, "Two"), (3, "Three"), (4, "Four")])
        self.assertEqual(direct.semantic_stream(), nested.semantic_stream())

    def test_overdepth_and_skipped_levels_rejected(self):
        for body in ("# One\n\n## Two\n\n### Three\n\n#### Four\n\n##### Five",
                     "## One\n\n### Two\n\n#### Three\n\n##### Four\n\n###### Five",
                     "# One\n\n## Two\n\n#### Four", "## One\n\n### Two\n\n##### Four"):
            with self.subTest(body=body), self.assertRaises(SourceError):
                self.write_source(body)

    def test_repeated_long_h4_titles_have_distinct_bookmarks_and_final_pages(self):
        title = "Long fourth heading " * 12
        parsed = self.write_source("# One\n\n## Two\n\n### Three\n\n#### " + title + "\n\nBody A.\n\n#### " + title + "\n\nBody B.")
        pdf = self.pdf(parsed.headings, [3, 3, 3, 4, 5])
        out = self.folder / "book.docx"
        build(self.source, out, pages_from_pdf=pdf)
        doc = Document(out)
        self.assertEqual(body_stream(doc), parsed.semantic_stream())
        h4 = [p for p in doc.paragraphs if p.style.name == "Heading 4"]
        toc4 = [p for p in doc.paragraphs if p.style.name == "Contents 4"]
        self.assertEqual(len(h4), 2)
        self.assertEqual([paragraph_xml_text(p) for p in toc4], [title.strip() + "\t4", title.strip() + "\t5"])
        names = [e.get(qn("w:name")) for e in doc.element.iter(qn("w:bookmarkStart"))]
        self.assertEqual(names, ["reader_heading_0001", "reader_heading_0002", "reader_heading_0003", "reader_heading_0004", "reader_heading_0005"])
        report = verify(self.source, out, pdf)
        self.assertTrue(report["ok"], report["errors"])
        h4[-1].text = "Altered fourth heading"
        doc.save(out)
        self.assertFalse(verify(self.source, out, pdf)["ok"])

    def test_default_h4_and_contents4_have_real_semantics_and_legible_roles(self):
        self.write_source("# One\n\n## Two\n\n### Three\n\n#### Four\n\nBody.")
        out = self.folder / "book.docx"
        build(self.source, out)
        doc = Document(out)
        h4 = doc.styles["Heading 4"]
        self.assertEqual(h4.font.size.pt, 12.5)
        self.assertTrue(h4.font.bold)
        self.assertEqual(str(h4.font.color.rgb), "C55A11")
        self.assertEqual(h4.element.find("w:pPr/w:outlineLvl", {"w": qn("w:p").split("}")[0][1:]}).get(qn("w:val")), "3")
        self.assertTrue(h4.paragraph_format.keep_with_next)
        self.assertEqual(h4.paragraph_format.first_line_indent, 0)
        c4 = doc.styles["Contents 4"]
        self.assertEqual(c4.font.size.pt, 10)
        self.assertEqual(c4.paragraph_format.left_indent.twips, 780)

    def legacy_template(self):
        legacy = self.folder / "legacy.docx"
        doc = Document(DEFAULT_TEMPLATE)
        for name in ("Heading 4", "Contents 4"):
            if name in doc.styles:
                doc.styles[name].delete()
        doc.save(legacy)
        return legacy

    def test_missing_optional_styles_derive_without_changing_existing_styles(self):
        legacy = self.legacy_template()
        before = Document(legacy)
        snapshot = {s.style_id: s.element.xml for s in before.styles}
        parsed = self.write_source("# One\n\n## Two\n\n### Three\n\n#### Four\n\nBody.")
        pdf = self.pdf(parsed.headings)
        out = self.folder / "book.docx"
        build(self.source, out, template=legacy, pages_from_pdf=pdf)
        after = Document(out)
        self.assertEqual(geometry(before), geometry(after))
        for style in after.styles:
            if style.style_id in snapshot:
                self.assertEqual(style.element.xml, snapshot[style.style_id])
        report = verify(self.source, out, pdf, legacy)
        self.assertTrue(report["ok"], report["errors"])

    def test_three_level_reference_distills_then_builds_four_levels(self):
        self.write_source("# One\n\n## Two\n\n### Three\n\nBody.")
        legacy = self.legacy_template()
        reference = self.folder / "reference.docx"
        build(self.source, reference, template=legacy)
        # A three-level reference need not contain H4/C4 paragraphs or definitions.
        doc = Document(reference)
        for name in ("Heading 4", "Contents 4"):
            if name in doc.styles:
                doc.styles[name].delete()
        doc.save(reference)
        original = reference.read_bytes()
        template = self.folder / "extracted.docx"
        distill(reference, template, None)
        self.assertEqual(reference.read_bytes(), original)
        parsed = self.write_source("# One\n\n## Two\n\n### Three\n\n#### Four\n\nBody.")
        pdf = self.pdf(parsed.headings)
        out = self.folder / "new.docx"
        build(self.source, out, template=template, pages_from_pdf=pdf)
        report = verify(self.source, out, pdf, template)
        self.assertTrue(report["ok"], report["errors"])

    def test_extraction_captures_actual_h4_local_typography(self):
        self.write_source("# One\n\n## Two\n\n### Three\n\n#### Four\n\nBody.")
        reference = self.folder / "reference.docx"
        build(self.source, reference)
        doc = Document(reference)
        p = next(p for p in doc.paragraphs if p.style.name == "Heading 4")
        p.runs[0].font.name = "Arial"
        p.runs[0].font.size = Pt(13)
        p.runs[0].font.bold = False
        p.runs[0].font.color.rgb = RGBColor.from_string("123456")
        p.paragraph_format.space_after = Pt(8)
        doc.save(reference)
        template = self.folder / "extracted.docx"
        evidence = distill(reference, template, None)
        self.assertEqual(evidence["roles"]["Heading 4"]["font"], "Arial")
        actual = Document(template).styles["Heading 4"]
        self.assertEqual(actual.font.size.pt, 13)
        self.assertFalse(actual.font.bold)
        self.assertEqual(str(actual.font.color.rgb), "123456")
        self.assertEqual(actual.paragraph_format.space_after.pt, 8)

    def test_sparse_chapter_tail_counts_body_lines_and_labels_book_end(self):
        parsed = self.write_source("# First\n\nText.\n\n# Last\n\nClosing text.")
        pdf = self.folder / "tails.pdf"
        with fitz.open() as doc:
            for _ in range(5):
                doc.new_page()
            for number in range(5):
                page = doc[number]
                page.insert_text((60, 25), "Header " + "x" * 90, fontsize=8)
                page.insert_text((60, 820), "Footer " + "x" * 90, fontsize=8)
            doc[0].insert_text((60, 100), "Book")
            doc[1].insert_text((60, 100), "Contents")
            doc[2].insert_text((60, 100), "First")
            for y in range(130, 430, 20):
                doc[2].insert_text((60, y), "Body " + "a" * 60, fontsize=9)
            # Many characters but fewer than three body lines still warns.
            doc[3].insert_text((60, 100), "a" * 105, fontsize=6)
            doc[3].insert_text((60, 120), "b" * 105, fontsize=6)
            doc[4].insert_text((60, 100), "Last")
            for y in (130, 150, 170):
                doc[4].insert_text((60, y), "Short closing text.", fontsize=9)
            doc.set_toc([[1, "First", 3], [1, "Last", 5]])
            doc.save(pdf)
        out = self.folder / "book.docx"
        build(self.source, out, pages_from_pdf=pdf)
        report = verify(self.source, out, pdf)
        self.assertTrue(report["ok"], report["errors"])
        tails = report["evidence"]["sparse_chapter_endings"]
        self.assertEqual([x["page"] for x in tails], [4, 5])
        self.assertEqual([x["classification"] for x in tails], ["chapter_ending", "book_ending"])
        self.assertEqual(tails[0]["body_lines"], 2)
        self.assertEqual(tails[0]["nonspace_body_characters"], 210)
        self.assertEqual(tails[1]["body_lines"], 3)
        self.assertEqual(tails[1]["nonspace_body_characters"], 51)
        self.assertTrue(report["warnings"])


if __name__ == "__main__":
    unittest.main()
