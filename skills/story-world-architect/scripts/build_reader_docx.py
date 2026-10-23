#!/usr/bin/env python3
"""Build a reader-facing world book from a deliberately small Markdown subset.

The first # heading is the cover title. Body headings are normalized from either
#/##/###/#### or ##/###/####/##### to Heading 1/2/3/4. Supported body blocks are paragraphs,
flat lists, and pipe tables.
Inline emphasis and single-backtick code are supported; unsupported Markdown
fails explicitly. Rendering and visual approval remain separate operations.
"""
from __future__ import annotations

import argparse
import json
import math
import re
import sys
from copy import deepcopy
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips

DEFAULT_TEMPLATE = Path(__file__).resolve().parent.parent / "assets" / "reader-template.docx"
ROLE_STYLES = (
    "Title", "Cover Subtitle", "Contents Title", "Contents 1", "Contents 2",
    "Contents 3", "Heading 1", "Heading 2", "Heading 3", "Reading Body",
    "Heading 4", "Contents 4",
)
OPTIONAL_STYLES = ("Heading 4", "Contents 4")
HEADING_STYLES = tuple(f"Heading {level}" for level in range(1, 5))
CONTENTS_STYLES = tuple(f"Contents {level}" for level in range(1, 5))
SLOTS = ("{{WORLD_TITLE}}", "{{WORLD_SUBTITLE}}", "{{TOC}}", "{{BODY}}")
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def prepare_reader_styles(document: Any) -> None:
    """Extend older clean templates without rewriting any existing style.

    Both build and verify use this same conditional derivation, so an older
    reference remains authoritative for its existing roles and geometry.
    """
    missing = [name for name in ROLE_STYLES if name not in OPTIONAL_STYLES and name not in document.styles]
    if missing:
        raise ValueError(f"Template is missing required styles: {missing}.")
    for name, base_name in (("Heading 4", "Heading 3"), ("Contents 4", "Contents 3")):
        if name in document.styles:
            continue
        source = document.styles[base_name]
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = source.base_style
        style.next_paragraph_style = source.next_paragraph_style
        for tag in ("pPr", "rPr"):
            node = source.element.find(qn("w:" + tag))
            if node is not None:
                style.element.append(deepcopy(node))
        ppr = style.element.get_or_add_pPr()
        if name == "Heading 4":
            style.quick_style = True
            style.font.size = Pt(12.5)
            style.font.bold = True
            style.font.italic = False
            style.font.color.rgb = RGBColor.from_string("C55A11")
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.first_line_indent = 0
            style.paragraph_format.space_before = Pt(10)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.page_break_before = False
            outline = ppr.find(qn("w:outlineLvl"))
            if outline is None:
                outline = OxmlElement("w:outlineLvl"); ppr.append(outline)
            outline.set(qn("w:val"), "3")
        else:
            style.font.size = Pt(10)
            style.paragraph_format.left_indent = Twips(780)
            style.paragraph_format.first_line_indent = 0
        # Match complex-script size to the new role instead of retaining H3/C3.
        size_cs = style.element.get_or_add_rPr().find(qn("w:szCs"))
        if size_cs is None:
            size_cs = OxmlElement("w:szCs"); style.element.rPr.append(size_cs)
        size_cs.set(qn("w:val"), "25" if name == "Heading 4" else "20")


class SourceError(ValueError):
    pass


@dataclass
class Span:
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False


@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    items: list[str] = field(default_factory=list)
    ordered: bool = False
    start: int = 1
    rows: list[list[str]] = field(default_factory=list)
    alignments: list[str] = field(default_factory=list)


@dataclass
class ParsedDocument:
    title: str
    blocks: list[Block]

    @property
    def headings(self) -> list[tuple[int, str]]:
        return [(b.level, plain_text(b.text)) for b in self.blocks if b.kind == "heading"]

    def semantic_stream(self) -> list[dict[str, Any]]:
        result: list[dict[str, Any]] = []
        for b in self.blocks:
            if b.kind == "heading":
                result.append({"kind": "heading", "level": b.level, "text": plain_text(b.text)})
            elif b.kind == "paragraph":
                result.append({"kind": "paragraph", "text": plain_text(b.text)})
            elif b.kind == "list":
                for offset, item in enumerate(b.items):
                    result.append({"kind": "list_item", "ordered": b.ordered,
                                   "number": b.start + offset if b.ordered else None,
                                   "text": plain_text(item)})
            elif b.kind == "table":
                result.append({"kind": "table", "rows": [[plain_text(c) for c in r] for r in b.rows]})
        return result


def inline_spans(text: str) -> list[Span]:
    """Parse balanced, simple emphasis/code without silently leaking Markdown."""
    if re.search(r"!?\[[^\]]*\]\s*(?:\(|\[)|\[\^[^\]]+\]|<[/!a-zA-Z][^>]*>", text):
        raise SourceError("Links, images, footnotes and HTML are not supported by this builder.")
    if re.search(r"\{\{[^{}]+\}\}", text):
        raise SourceError("Source contains a template placeholder; provide finished reader text.")
    if re.search(r"(?<!\\)~~|(?<!\\)\$\S[^$]*?\$", text):
        raise SourceError("Strikethrough and inline mathematics are not supported.")
    result: list[Span] = []

    def walk(value: str, bold: bool = False, italic: bool = False) -> None:
        buf: list[str] = []

        def flush() -> None:
            if buf:
                result.append(Span("".join(buf), bold, italic))
                buf.clear()

        i = 0
        while i < len(value):
            char = value[i]
            if char == "\\":
                if i + 1 >= len(value) or value[i + 1] not in r"\`*_{}[]()#+-.!|>$~":
                    raise SourceError("Unsupported backslash escape; use ordinary text or escape Markdown punctuation.")
                buf.append(value[i + 1]); i += 2; continue
            if char == "`":
                if value.startswith("``", i):
                    raise SourceError("Only single-backtick inline code is supported.")
                end = value.find("`", i + 1)
                if end < 0 or end == i + 1:
                    raise SourceError("Unbalanced or empty inline code.")
                flush(); result.append(Span(value[i + 1:end], bold, italic, True)); i = end + 1; continue
            if char in "*_":
                # Underscores within ordinary identifiers remain literal.
                if char == "_" and i and i + 1 < len(value) and value[i - 1].isalnum() and value[i + 1].isalnum():
                    buf.append(char); i += 1; continue
                marker = char * (2 if value.startswith(char * 2, i) else 1)
                if value.startswith(char * 3, i):
                    raise SourceError("Triple emphasis is not supported; use simple bold or italic.")
                end = i + len(marker)
                while True:
                    end = value.find(marker, end)
                    if end < 0:
                        raise SourceError("Unbalanced emphasis; escape literal asterisks or underscores.")
                    if end == 0 or value[end - 1] != "\\":
                        break
                    end += len(marker)
                inner = value[i + len(marker):end]
                if not inner or inner.isspace():
                    raise SourceError("Empty emphasis is not supported.")
                flush(); walk(inner, bold or len(marker) == 2, italic or len(marker) == 1)
                i = end + len(marker); continue
            buf.append(char); i += 1
        flush()

    walk(text)
    return result


def plain_text(text: str) -> str:
    return "".join(s.text for s in inline_spans(text))


def normalized_heading(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def table_cells(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    return [c.strip() for c in re.split(r"(?<!\\)\|", line)]


def table_separator(line: str) -> list[str] | None:
    if "|" not in line:
        return None
    cells = table_cells(line)
    if not cells or not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
        return None
    return ["center" if c.startswith(":") and c.endswith(":") else "right" if c.endswith(":") else "left" for c in cells]


def parse_markdown(source: Path | str) -> ParsedDocument:
    raw = Path(source).read_text(encoding="utf-8-sig")
    lines = raw.splitlines()
    nonblank = [i for i, line in enumerate(lines) if line.strip()]
    if not nonblank or not re.fullmatch(r"#\s+\S.*", lines[nonblank[0]].strip()):
        raise SourceError("The first nonblank line must be '# Book title'.")
    first = nonblank[0]
    title = re.sub(r"\s+#+\s*$", "", lines[first].strip()[2:]).strip()
    title = plain_text(title)
    blocks: list[Block] = []
    raw_levels = [len(match[1]) for line in lines[first + 1:]
                  if (match := re.fullmatch(r"(#{1,6})\s+(\S.*)", line.strip()))]
    heading_base = min(raw_levels, default=1)
    if heading_base not in (1, 2):
        raise SourceError("Body headings must start with # or ##; up to four normalized levels are supported.")
    previous_level = 0
    i = first + 1
    list_re = re.compile(r"^(?:([-+*])|(\d+)[.)])\s+(\S.*)$")

    def reject_special(line: str, number: int) -> None:
        if line.startswith("\t") or len(line) - len(line.lstrip(" ")) >= 4:
            raise SourceError(f"Line {number}: indented code and nested lists are not supported.")
        s = line.strip()
        if re.match(r"(?:```|~~~|>|!\[|\[\^|\$\$)", s) or re.fullmatch(r"(?:[-*_]\s*){3,}", s):
            raise SourceError(f"Line {number}: unsupported code, quote, image, footnote, math or horizontal rule.")
        if re.match(r"<[/!a-zA-Z]", s) or re.match(r"\[[^\]]+\]:", s):
            raise SourceError(f"Line {number}: HTML and reference links are not supported.")
        if (match := re.match(r"(#+)\s", s)) and len(match[1]) > heading_base + 3:
            raise SourceError(f"Line {number}: only four normalized body heading levels are supported.")
        if line.endswith("  ") or line.endswith("\\"):
            raise SourceError(f"Line {number}: explicit Markdown line breaks are not supported.")

    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1; continue
        reject_special(line, i + 1)
        s = line.strip()
        heading = re.fullmatch(r"(#{1,6})\s+(\S.*)", s)
        if heading:
            level = len(heading[1]) - heading_base + 1
            if level > previous_level + 1:
                raise SourceError(f"Line {i + 1}: heading hierarchy jumps to level {level}.")
            previous_level = level
            text = re.sub(r"\s+#+\s*$", "", heading[2]).strip()
            plain_text(text); blocks.append(Block("heading", text=text, level=level)); i += 1; continue
        alignment = table_separator(lines[i + 1]) if i + 1 < len(lines) else None
        if alignment is not None:
            header = table_cells(s)
            if len(header) != len(alignment):
                raise SourceError(f"Line {i + 1}: table header/separator column counts differ.")
            rows = [header]; i += 2
            while i < len(lines) and lines[i].strip() and "|" in lines[i]:
                reject_special(lines[i], i + 1)
                row = table_cells(lines[i])
                if len(row) != len(header):
                    raise SourceError(f"Line {i + 1}: table column count differs from header.")
                rows.append(row); i += 1
            for row in rows:
                for cell in row:
                    plain_text(cell)
            blocks.append(Block("table", rows=rows, alignments=alignment)); continue
        item = list_re.fullmatch(s)
        if item:
            ordered = item[2] is not None
            start = int(item[2]) if ordered else 1
            if start < 1:
                raise SourceError(f"Line {i + 1}: ordered lists must start at a positive number.")
            items: list[str] = []
            while i < len(lines):
                current = lines[i]
                if not current.strip():
                    # One blank between simple items is accepted.
                    if i + 1 < len(lines) and list_re.fullmatch(lines[i + 1].strip()):
                        i += 1; continue
                    break
                reject_special(current, i + 1)
                match = list_re.fullmatch(current.strip())
                if not match:
                    if not re.match(r"#{1,6}\s", current.strip()) and not (
                        i + 1 < len(lines) and table_separator(lines[i + 1]) is not None
                    ):
                        raise SourceError(f"Line {i + 1}: use one line per simple list item and a blank line before the next paragraph.")
                    break
                if (match[2] is not None) != ordered:
                    break
                if current != current.lstrip():
                    raise SourceError(f"Line {i + 1}: nested or indented list items are not supported.")
                if ordered and int(match[2]) != start + len(items):
                    raise SourceError(f"Line {i + 1}: ordered list numbers must be sequential.")
                if re.match(r"\[[ xX]\]\s", match[3]):
                    raise SourceError(f"Line {i + 1}: task lists are not supported.")
                plain_text(match[3]); items.append(match[3]); i += 1
            blocks.append(Block("list", items=items, ordered=ordered, start=start)); continue
        if s.startswith("|") or table_separator(s):
            raise SourceError(f"Line {i + 1}: malformed table or unsupported table layout.")
        paragraph = [s]; i += 1
        while i < len(lines) and lines[i].strip():
            next_s = lines[i].strip()
            if re.match(r"#{1,6}\s", next_s) or list_re.fullmatch(next_s):
                break
            if i + 1 < len(lines) and table_separator(lines[i + 1]) is not None:
                break
            reject_special(lines[i], i + 1)
            if next_s.startswith("|") or re.fullmatch(r"[=-]{3,}", next_s):
                raise SourceError(f"Line {i + 1}: malformed table or setext heading is not supported.")
            paragraph.append(next_s); i += 1
        text = " ".join(paragraph)
        plain_text(text); blocks.append(Block("paragraph", text=text))
    return ParsedDocument(title, blocks)


def _extract_pdf_outline(pdf: Path | str) -> tuple[list[list[Any]], int]:
    try:
        import fitz
    except ImportError:
        fitz = None  # type: ignore[assignment]
    if fitz is not None:
        with fitz.open(pdf) as document:
            if document.needs_pass:
                raise ValueError("Encrypted PDFs must be unlocked before reading outlines.")
            return document.get_toc(simple=True), len(document)
    try:
        import pypdf
    except ImportError as exc:
        raise RuntimeError("Neither PyMuPDF (fitz) nor pypdf is installed.") from exc
    reader = pypdf.PdfReader(str(pdf))
    if reader.is_encrypted:
        raise ValueError("Encrypted PDFs must be unlocked before reading outlines.")

    def walk(items: list[Any], level: int = 1) -> list[list[Any]]:
        entries: list[list[Any]] = []
        for item in items:
            if isinstance(item, list):
                entries.extend(walk(item, level + 1))
            else:
                page_idx = reader.get_destination_page_number(item)
                page_num = (page_idx + 1) if page_idx is not None else -1
                entries.append([level, item.title or "", page_num])
        return entries

    return walk(reader.outline), len(reader.pages)


def pdf_heading_pages(pdf: Path | str, headings: list[tuple[int, str]]) -> list[int]:
    outline, total_pages = _extract_pdf_outline(pdf)
    if len(outline) != len(headings):
        raise ValueError(f"PDF outline has {len(outline)} entries; source has {len(headings)} headings.")
    pages: list[int] = []
    for index, ((level, title), entry) in enumerate(zip(headings, outline), 1):
        if level != entry[0] or normalized_heading(title) != normalized_heading(entry[1]):
            raise ValueError(f"PDF outline entry {index} differs: expected ({level}, {title!r}); got {entry[:2]!r}.")
        if not 1 <= entry[2] <= total_pages:
            raise ValueError(f"PDF outline entry {index} has an invalid page: {entry[2]}.")
        pages.append(entry[2])
    return pages


def paragraph_xml_text(paragraph: Any) -> str:
    result: list[str] = []
    for child in paragraph._p:
        if child.tag == qn("w:pPr"):
            continue
        for element in child.iter():
            if element.tag == qn("w:t"):
                result.append(element.text or "")
            elif element.tag == qn("w:tab"):
                result.append("\t")
            elif element.tag == qn("w:br"):
                result.append("\n")
    return "".join(result)


def fill_paragraph(paragraph: Any, text: str) -> None:
    paragraph.clear()
    for span in inline_spans(text):
        run = paragraph.add_run(span.text)
        if span.bold:
            run.bold = True
        if span.italic:
            run.italic = True
        if span.code:
            shading = OxmlElement("w:shd"); shading.set(qn("w:fill"), "F1F3F5")
            run._r.get_or_add_rPr().append(shading)


def make_numbering(document: Any, ordered: bool, start: int) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum"); abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType"); multi.set(qn("w:val"), "singleLevel"); abstract.append(multi)
    level = OxmlElement("w:lvl"); level.set(qn("w:ilvl"), "0")
    for name, val in (("start", str(start)), ("numFmt", "decimal" if ordered else "bullet"),
                      ("lvlText", "%1." if ordered else "•"), ("lvlJc", "left")):
        child = OxmlElement("w:" + name); child.set(qn("w:val"), val); level.append(child)
    abstract.append(level); numbering.append(abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(num_id))
    reference = OxmlElement("w:abstractNumId"); reference.set(qn("w:val"), str(abstract_id)); num.append(reference)
    numbering.append(num)
    return num_id


def add_bookmark(paragraph: Any, name: str, identifier: int) -> None:
    start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"), str(identifier)); start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), str(identifier))
    ppr = paragraph._p.find(qn("w:pPr"))
    paragraph._p.insert(1 if ppr is not None else 0, start); paragraph._p.append(end)


def add_toc_entry(anchor: Any, level: int, title: str, page: int, target: str, width: int) -> None:
    paragraph = anchor.insert_paragraph_before(style=f"Contents {level}")
    paragraph.paragraph_format.tab_stops.clear_all()
    paragraph.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    hyperlink = OxmlElement("w:hyperlink"); hyperlink.set(qn("w:anchor"), target); hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r"); text = OxmlElement("w:t"); text.text = title; run.append(text); hyperlink.append(run)
    paragraph._p.append(hyperlink); paragraph.add_run("\t" + str(page))


def build(source: Path, out: Path, template: Path = DEFAULT_TEMPLATE,
          subtitle: str = "世界设定集", pages_from_pdf: Path | None = None) -> dict[str, Any]:
    parsed = parse_markdown(source)
    if not parsed.headings:
        raise SourceError("A world book needs at least one body heading for its contents.")
    if out.resolve() in (source.resolve(), template.resolve()):
        raise ValueError("Output must not overwrite the source Markdown or the template.")
    document = Document(template)
    prepare_reader_styles(document)
    if len(document.sections) != 1:
        raise ValueError("This basic builder requires a single-section, text-only reader template.")
    slots: dict[str, Any] = {}
    for token in SLOTS:
        matches = [p for p in document.paragraphs if paragraph_xml_text(p) == token]
        if len(matches) != 1:
            raise ValueError(f"Template must contain exactly one body paragraph with {token}; found {len(matches)}.")
        slots[token] = matches[0]
    pages = pdf_heading_pages(pages_from_pdf, parsed.headings) if pages_from_pdf else [1] * len(parsed.headings)
    # Slots have already inherited the template's cover roles. Preserve pPr.
    slots["{{WORLD_TITLE}}"].clear(); slots["{{WORLD_TITLE}}"].add_run(parsed.title)
    fill_paragraph(slots["{{WORLD_SUBTITLE}}"], subtitle)
    for section in document.sections:
        for container in (section.header, section.first_page_header, section.even_page_header):
            for paragraph in container.paragraphs:
                old = paragraph_xml_text(paragraph)
                if "{{WORLD_TITLE}}" in old:
                    # Retain the template's local header typography.
                    rpr = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
                    paragraph.clear()
                    run = paragraph.add_run(old.replace("{{WORLD_TITLE}}", parsed.title))
                    if rpr is not None:
                        run._r.insert(0, rpr)
    body_anchor = slots["{{BODY}}"]
    fill_paragraph(slots["{{TOC}}"], "目录")
    toc_anchor = body_anchor.insert_paragraph_before()
    section = document.sections[0]
    width = section.page_width - section.left_margin - section.right_margin
    bookmark_ids = [int(x.get(qn("w:id"))) for x in document.element.iter(qn("w:bookmarkStart"))]
    base_id = max(bookmark_ids, default=0) + 1
    heading_index = 0
    for block_index, block in enumerate(parsed.blocks):
        if block.kind in ("paragraph", "heading"):
            style = f"Heading {block.level}" if block.kind == "heading" else "Reading Body"
            paragraph = body_anchor.insert_paragraph_before(style=style)
            fill_paragraph(paragraph, block.text)
            if block_index == 0:
                paragraph.paragraph_format.page_break_before = True
            if block.kind == "heading":
                heading_index += 1
                target = f"reader_heading_{heading_index:04d}"
                add_bookmark(paragraph, target, base_id + heading_index - 1)
                add_toc_entry(toc_anchor, block.level, plain_text(block.text), pages[heading_index - 1], target, width)
        elif block.kind == "list":
            number_id = make_numbering(document, block.ordered, block.start)
            for item in block.items:
                paragraph = body_anchor.insert_paragraph_before(style="Reading Body")
                fill_paragraph(paragraph, item)
                paragraph.paragraph_format.left_indent = Inches(0.28)
                paragraph.paragraph_format.first_line_indent = Inches(-0.18)
                paragraph.paragraph_format.space_after = Pt(3)
                numpr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                numpr.get_or_add_ilvl().val = 0; numpr.get_or_add_numId().val = number_id
        elif block.kind == "table":
            table = document.add_table(rows=len(block.rows), cols=len(block.rows[0]))
            if "Table Grid" in document.styles:
                table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
            tblpr = table._tbl.tblPr
            borders = tblpr.find(qn("w:tblBorders"))
            if borders is not None:
                tblpr.remove(borders)
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                node = OxmlElement("w:" + edge)
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    node.set(qn("w:" + key), value)
                borders.append(node)
            tblpr.append(borders)
            margins = OxmlElement("w:tblCellMar")
            for edge in ("top", "bottom", "left", "right"):
                node = OxmlElement("w:" + edge); node.set(qn("w:w"), "90"); node.set(qn("w:type"), "dxa"); margins.append(node)
            tblpr.append(margins)
            # Give descriptive columns more space without allowing a single long
            # cell to dominate. This is only a starting layout; render review is required.
            weights = [math.sqrt(max(4, max(sum(1 if ord(c) > 255 else 0.55 for c in plain_text(row[ci]))
                                          for row in block.rows))) for ci in range(len(block.rows[0]))]
            widths = [int(width * weight / sum(weights)) for weight in weights]
            for ci, col in enumerate(table.columns):
                col.width = widths[ci]
            for ri, row in enumerate(block.rows):
                for ci, text in enumerate(row):
                    cell = table.cell(ri, ci); cell.width = widths[ci]
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    paragraph = cell.paragraphs[0]; paragraph.style = document.styles["Reading Body"]
                    paragraph.paragraph_format.first_line_indent = 0
                    paragraph.paragraph_format.left_indent = 0; paragraph.paragraph_format.right_indent = 0
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.1
                    from docx.enum.text import WD_ALIGN_PARAGRAPH
                    paragraph.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                                           "center": WD_ALIGN_PARAGRAPH.CENTER}[block.alignments[ci]]
                    fill_paragraph(paragraph, text)
                    if ri == 0:
                        shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "F1F3F5"); cell._tc.get_or_add_tcPr().append(shade)
                        for run in paragraph.runs:
                            run.bold = True
                trpr = table.rows[ri]._tr.get_or_add_trPr()
                if ri == 0:
                    repeat = OxmlElement("w:tblHeader"); trpr.append(repeat)
            body_anchor._p.addprevious(table._tbl)
    for paragraph in (toc_anchor, body_anchor):
        paragraph._p.getparent().remove(paragraph._p)
    for element in document.element.iter(qn("w:t")):
        if element.text and re.search(r"\{\{[^{}]+\}\}", element.text):
            raise ValueError(f"An unfilled template slot remains: {element.text!r}.")
    out.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = parsed.title
    document.core_properties.subject = plain_text(subtitle)
    document.save(out)
    return {"output": str(out.resolve()), "title": parsed.title, "headings": len(parsed.headings),
            "body_blocks": len(parsed.blocks), "toc_pages_from_pdf": str(pages_from_pdf) if pages_from_pdf else None,
            "requires_final_render": True}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--subtitle", default="世界设定集")
    parser.add_argument("--pages-from-pdf", type=Path)
    args = parser.parse_args()
    try:
        result = build(args.source, args.out, args.template, args.subtitle, args.pages_from_pdf)
    except (OSError, ValueError, RuntimeError, KeyError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr); return 2
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
