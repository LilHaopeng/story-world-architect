#!/usr/bin/env python3
"""Distill a single-section text worldbook into a content-free reusable template."""
from __future__ import annotations

import argparse
from collections import Counter
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree


ROLES = ("Title", "Cover Subtitle", "Contents Title", "Contents 1", "Contents 2", "Contents 3",
         "Heading 1", "Heading 2", "Heading 3", "Reading Body")


def sanitize_template_metadata(path: Path) -> None:
    """Clear authoring fingerprints and stale statistics in an extracted template.

    This is not a content-redaction or general DOCX privacy guarantee.
    The extractor removes custom properties and original content separately.
    """
    buffer = BytesIO()
    word_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    with ZipFile(path) as original, ZipFile(buffer, "w") as cleaned:
        for entry in original.infolist():
            data = original.read(entry.filename)
            if entry.filename == "docProps/app.xml":
                data = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"><Application>story-world-architect</Application></Properties>'
            elif entry.filename == "docProps/core.xml":
                data = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"/>'
            elif entry.filename.startswith("word/") and entry.filename.endswith(".xml"):
                root = etree.fromstring(data)
                for element in root.iter():
                    for attribute in list(element.attrib):
                        if attribute.startswith("{" + word_ns + "}rsid"):
                            del element.attrib[attribute]
                for element in list(root.iter()):
                    if element.tag in ("{" + word_ns + "}rsid", "{" + word_ns + "}rsids") and element.getparent() is not None:
                        element.getparent().remove(element)
                data = etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)
            cleaned.writestr(entry, data)
    path.write_bytes(buffer.getvalue())


def inherited_font(style, key):
    while style is not None:
        value = getattr(style.font, key)
        if value is not None:
            return value
        style = style.base_style
    return None


def effective(run, style, key):
    value = getattr(run.font, key)
    return value if value is not None else inherited_font(style, key)


def font_slots(style, font_name):
    style.font.name = font_name
    rfonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme", "csTheme"):
        rfonts.attrib.pop(qn("w:" + key), None)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn("w:" + key), font_name)


def distill(reference: Path, output: Path, evidence: Path | None):
    if reference.resolve() == output.resolve():
        raise ValueError("The retained reference must not be overwritten.")
    original_hash = sha256(reference.read_bytes()).hexdigest()
    doc = Document(reference)
    if len(doc.sections) != 1 or doc.tables or doc.inline_shapes:
        raise ValueError("Only single-section text references are supported; use the documents fidelity workflow for complex templates.")
    if doc._element.xpath(".//w:drawing | .//w:pict | .//w:sdt | .//w:txbxContent | .//w:ins | .//w:del | .//w:commentRangeStart"):
        raise ValueError("Reference has unsupported drawings, controls or revisions; inspect it with the documents workflow.")
    for name in ROLES:
        if name not in doc.styles or not any(p.style.name == name for p in doc.paragraphs):
            raise ValueError(f"Missing named reference role: {name}")

    first_title = next(i for i, p in enumerate(doc.paragraphs) if p.style.name == "Title")
    spacers = [deepcopy(p._p) for p in doc.paragraphs[:first_title] if not p.text.strip()]
    roles = {}
    for name in ROLES:
        style = doc.styles[name]
        paragraphs = [p for p in doc.paragraphs if p.style.name == name]
        if paragraphs[0]._p.pPr is not None:
            target = style._element.get_or_add_pPr()
            for item in paragraphs[0]._p.pPr:
                if item.tag == qn("w:pStyle"):
                    continue
                for old in list(target):
                    if old.tag == item.tag:
                        target.remove(old)
                target.append(deepcopy(item))
        formats = []
        for p in paragraphs:
            for run in p.runs:
                if not run.text.strip():
                    continue
                font = effective(run, style, "name")
                size = effective(run, style, "size")
                bold = effective(run, style, "bold")
                color = run.font.color.rgb
                if color is None:
                    color = style.font.color.rgb
                formats.append((font, size.pt if size else None, bold, str(color) if color else "000000"))
        if not formats:
            formats = [(style.font.name, style.font.size.pt if style.font.size else None, style.font.bold, "000000")]
        font, size, bold, color = Counter(formats).most_common(1)[0][0]
        if font is None or size is None:
            raise ValueError(f"Cannot establish effective font and size for {name}")
        font_slots(style, font)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        roles[name] = {
            "font": font, "pt": size, "bold": bold, "color": color,
            "paragraph_xml": etree.tostring(style._element.pPr, encoding="unicode") if style._element.pPr is not None else None,
        }

    # Source text must not survive in opaque metadata or unused notes.
    for rid, rel in list(doc.part.rels.items()):
        suffix = rel.reltype.rsplit("/", 1)[-1]
        if suffix in {"footnotes", "endnotes"}:
            root = etree.fromstring(rel.target_part.blob)
            if any(t.text and t.text.strip() for t in root.iter(qn("w:t"))):
                raise ValueError("Meaningful notes require explicit template distillation.")
        if suffix in {"customXml", "footnotes", "endnotes", "comments"}:
            doc.part.drop_rel(rid)
    for rid, rel in list(doc.part.package.rels.items()):
        if rel.reltype.rsplit("/", 1)[-1] in {"custom-properties", "thumbnail"}:
            del doc.part.package.rels[rid]
    for item in doc.settings._element.xpath(".//w:footnotePr | .//w:endnotePr"):
        item.getparent().remove(item)
    for key in ("title", "subject", "author", "last_modified_by", "keywords", "comments", "category", "identifier", "content_status"):
        setattr(doc.core_properties, key, "")

    body = doc._element.body
    for item in list(body):
        if item.tag != qn("w:sectPr"):
            body.remove(item)
    for spacer in spacers:
        body.insert(len(body) - 1, spacer)
    doc.add_paragraph("{{WORLD_TITLE}}", "Title")
    doc.add_paragraph("{{WORLD_SUBTITLE}}", "Cover Subtitle")
    doc.add_paragraph("{{TOC}}", "Contents Title")
    doc.add_paragraph("{{BODY}}", "Reading Body")
    for section in doc.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            for p in header.paragraphs:
                if not p.text.strip():
                    continue
                run_properties = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
                p.clear()
                run = p.add_run("{{WORLD_TITLE}}")
                if run_properties is not None:
                    run._r.insert(0, run_properties)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    sanitize_template_metadata(output)
    assert sha256(reference.read_bytes()).hexdigest() == original_hash
    with ZipFile(output) as archive:
        if archive.testzip() is not None:
            raise ValueError("Invalid generated package")
    section = doc.sections[0]
    result = {
        "reference": str(reference.resolve()), "reference_sha256": original_hash,
        "template_sha256": sha256(output.read_bytes()).hexdigest(), "roles": roles,
        "page_twips": {key: getattr(section, key).twips for key in (
            "page_width", "page_height", "top_margin", "bottom_margin", "left_margin", "right_margin", "header_distance", "footer_distance")},
        "different_first_page": section.different_first_page_header_footer,
        "source_unchanged": True,
    }
    if evidence:
        evidence.parent.mkdir(parents=True, exist_ok=True)
        evidence.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--reference", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--evidence", type=Path)
    args = parser.parse_args()
    result = distill(args.reference, args.out, args.evidence)
    print(json.dumps({"template": str(args.out.resolve()), "source_unchanged": result["source_unchanged"], "roles": len(result["roles"])}, ensure_ascii=True))


if __name__ == "__main__":
    main()
