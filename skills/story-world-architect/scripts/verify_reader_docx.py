#!/usr/bin/env python3
"""Verify measurable structure of a reader DOCX and its final rendered PDF.

This is not a semantic or visual approval. Warnings (including sparse pages)
require human review and do not automatically make a book fail.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

try:
    import fitz
except ImportError:
    fitz = None  # type: ignore[assignment]


class _FallbackRect:
    def __init__(self, width: float, height: float) -> None:
        self.width = float(width)
        self.height = float(height)


class _FallbackPage:
    def __init__(self, width: float, height: float, text: str, lines: list[dict[str, Any]]) -> None:
        self.rect = _FallbackRect(width, height)
        self._text = text
        self._lines = lines
        self._spans: list[dict[str, Any]] = []

    def get_text(self, mode: str = "text", sort: bool = False) -> Any:
        if mode == "dict":
            if not sort:
                return {"blocks": [{"lines": [{"spans": self._spans}]}]}
            lines = sorted(self._lines, key=lambda l: (l["bbox"][1], l["bbox"][0])) if sort else self._lines
            return {"blocks": [{"lines": lines}]}
        return self._text


class _FallbackDocument:
    def __init__(self, pdf: Path) -> None:
        self.pdf = pdf
        self._pages: list[_FallbackPage] = []

    def __enter__(self) -> _FallbackDocument:
        try:
            import pdfplumber
        except ImportError as exc:
            raise RuntimeError("Layout verification requires PyMuPDF (fitz) or pdfplumber; pypdf alone cannot measure text bounds.") from exc
        with pdfplumber.open(str(self.pdf)) as doc:
            for page in doc.pages:
                # Preserve real PDF character bounds and mixed-font lines. Never
                # invent positions from plain extracted text to claim layout QA.
                spans = [{"text": char.get("text", ""), "font": char.get("fontname", "unknown"),
                          "bbox": tuple(float(char[key]) for key in ("x0", "top", "x1", "bottom"))}
                         for char in page.chars]
                raw_lines = page.extract_text_lines() or []
                lines = [{"bbox": tuple(float(line[key]) for key in ("x0", "top", "x1", "bottom")),
                          "spans": [{"text": line.get("text", ""), "font": "unknown",
                                     "bbox": tuple(float(line[key]) for key in ("x0", "top", "x1", "bottom"))}]}
                         for line in raw_lines]
                # Unsorted access is for glyph bounds/font inventory; sorted
                # access is for the rendered-line chapter-ending heuristic.
                self._pages.append(_FallbackPage(float(page.width), float(page.height), page.extract_text() or "", lines))
                self._pages[-1]._spans = spans
        return self

    def __exit__(self, exc_type: Any, exc: Any, tb: Any) -> None:
        pass

    def __len__(self) -> int:
        return len(self._pages)

    def __getitem__(self, index: int) -> _FallbackPage:
        return self._pages[index]

    def __iter__(self) -> Any:
        return iter(self._pages)


def _open_pdf(pdf: Path) -> Any:
    if fitz is not None:
        return fitz.open(pdf)
    return _FallbackDocument(pdf)


from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from build_reader_docx import (CONTENTS_STYLES, DEFAULT_TEMPLATE, HEADING_STYLES,
                               NS, ROLE_STYLES, parse_markdown, prepare_reader_styles,
                               paragraph_xml_text, pdf_heading_pages)

FONT_SLOTS = ("ascii", "hAnsi", "eastAsia", "cs")
THEME_SLOTS = ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme", "csTheme")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def style_fonts(document: Any, style: Any) -> dict[str, str | None]:
    result: dict[str, str | None] = {slot: None for slot in FONT_SLOTS}
    seen: set[str] = set()
    current = style
    while current is not None and current.style_id not in seen:
        seen.add(current.style_id)
        fonts = current.element.find("w:rPr/w:rFonts", NS)
        if fonts is not None:
            for slot in FONT_SLOTS:
                result[slot] = result[slot] or fonts.get(qn("w:" + slot))
        current = current.base_style
    defaults = document.styles.element.find("w:docDefaults/w:rPrDefault/w:rPr/w:rFonts", NS)
    if defaults is not None:
        for slot in FONT_SLOTS:
            result[slot] = result[slot] or defaults.get(qn("w:" + slot))
    return result


def geometry(document: Any) -> list[dict[str, Any]]:
    fields = ("page_width", "page_height", "left_margin", "right_margin", "top_margin", "bottom_margin",
              "header_distance", "footer_distance", "gutter", "orientation", "start_type")
    return [{field: int(value) if (value := getattr(section, field)) is not None else None for field in fields}
            for section in document.sections]


def numbering_definition(document: Any, number_id: int) -> tuple[bool, int]:
    root = document.part.numbering_part.element
    nums = [n for n in root.findall(qn("w:num")) if n.get(qn("w:numId")) == str(number_id)]
    if len(nums) != 1:
        raise ValueError(f"Numbering ID {number_id} is absent or duplicated.")
    abstract_id = nums[0].find(qn("w:abstractNumId")).get(qn("w:val"))
    abstracts = [a for a in root.findall(qn("w:abstractNum")) if a.get(qn("w:abstractNumId")) == abstract_id]
    if len(abstracts) != 1:
        raise ValueError(f"Abstract numbering {abstract_id} is absent or duplicated.")
    levels = [x for x in abstracts[0].findall(qn("w:lvl")) if x.get(qn("w:ilvl")) == "0"]
    if len(levels) != 1:
        raise ValueError("Only single-level numbering is supported.")
    fmt = levels[0].find(qn("w:numFmt")).get(qn("w:val"))
    if fmt not in ("decimal", "bullet"):
        raise ValueError(f"Unsupported list format {fmt}.")
    start_node = levels[0].find(qn("w:start"))
    start = int(start_node.get(qn("w:val"))) if start_node is not None else 1
    for override in nums[0].findall(qn("w:lvlOverride")):
        if override.get(qn("w:ilvl")) == "0":
            value = override.find(qn("w:startOverride"))
            if value is not None:
                start = int(value.get(qn("w:val")))
    return fmt == "decimal", start


def body_stream(document: Any) -> list[dict[str, Any]]:
    stream: list[dict[str, Any]] = []
    counters: dict[int, int] = {}
    for child in document.element.body:
        if child.tag == qn("w:tbl"):
            table = Table(child, document._body)
            stream.append({"kind": "table", "rows": [[cell.text for cell in row.cells] for row in table.rows]})
        elif child.tag == qn("w:p"):
            paragraph = Paragraph(child, document._body)
            name = paragraph.style.name
            text = paragraph_xml_text(paragraph)
            if name in HEADING_STYLES:
                stream.append({"kind": "heading", "level": int(name[-1]), "text": text})
            elif name == "Reading Body":
                numpr = paragraph._p.find("w:pPr/w:numPr", NS)
                if numpr is None:
                    stream.append({"kind": "paragraph", "text": text})
                else:
                    level = numpr.find(qn("w:ilvl"))
                    if level is not None and level.get(qn("w:val")) != "0":
                        raise ValueError("Nested lists are not supported.")
                    number_id = int(numpr.find(qn("w:numId")).get(qn("w:val")))
                    ordered, start = numbering_definition(document, number_id)
                    value = counters.get(number_id, start - 1) + 1; counters[number_id] = value
                    stream.append({"kind": "list_item", "ordered": ordered,
                                   "number": value if ordered else None, "text": text})
    return stream


def chapter_ending_warnings(rendered: Any, reference: Any,
                            headings: list[tuple[int, str]], pages: list[int]) -> list[dict[str, Any]]:
    """Measure rendered body lines at chapter endings, not words or paragraphs.

    The template's body region removes headers/footers. Outline titles remove
    heading lines (including wrapped titles). Pages preceding the first body
    heading are cover/contents, outside this check. This is a review heuristic.
    """
    if len(pages) != len(headings) or not pages:
        return []
    section = reference.sections[0]
    measured: dict[int, tuple[int, int]] = {}
    for number in range(pages[0], len(rendered) + 1):
        page = rendered[number - 1]
        scale = page.rect.height / section.page_height.pt
        top = section.top_margin.pt * scale
        bottom = page.rect.height - section.bottom_margin.pt * scale
        lines = ["".join(span["text"] for span in line["spans"]).strip()
                 for block in page.get_text("dict", sort=True)["blocks"]
                 for line in block.get("lines", [])
                 if line["bbox"][1] >= top - 2 and line["bbox"][3] <= bottom + 2]
        lines = [text for text in lines if text]
        excluded: set[int] = set()
        # Match complete line groups; never remove an arbitrary body substring.
        for (_, title), heading_page in zip(headings, pages):
            if heading_page != number:
                continue
            wanted = re.sub(r"\s+", "", title)
            for start in range(len(lines)):
                if start in excluded:
                    continue
                joined = ""
                for end in range(start, len(lines)):
                    if end in excluded:
                        break
                    joined += re.sub(r"\s+", "", lines[end])
                    if joined == wanted:
                        excluded.update(range(start, end + 1)); break
                    if not wanted.startswith(joined):
                        break
                else:
                    continue
                if start in excluded:
                    break
        body = [text for index, text in enumerate(lines) if index not in excluded]
        measured[number] = (sum(len(re.sub(r"\s+", "", text)) for text in body), len(body))
    chapters = [(title, page) for (level, title), page in zip(headings, pages) if level == 1]
    nonempty = [page for page, (chars, _) in measured.items() if chars]
    final_page = max(nonempty, default=pages[-1])
    result: list[dict[str, Any]] = []
    for index, (title, start) in enumerate(chapters):
        is_end = index == len(chapters) - 1
        end = final_page if is_end else max(start, chapters[index + 1][1] - 1)
        chars, line_count = measured.get(end, (0, 0))
        if chars < 200 or line_count < 3:
            result.append({"chapter": title, "page": end,
                           "classification": "book_ending" if is_end else "chapter_ending",
                           "nonspace_body_characters": chars, "body_lines": line_count,
                           "review_only": True})
    return result


def verify(source: Path, docx: Path, pdf: Path, template: Path = DEFAULT_TEMPLATE) -> dict[str, Any]:
    report: dict[str, Any] = {"ok": False, "errors": [], "warnings": [], "checks": {}, "evidence": {},
                             "validation_scope": "Structural, textual and measurable layout checks only; semantic and visual review remain required."}
    errors: list[str] = report["errors"]
    warnings: list[str] = report["warnings"]
    parsed = parse_markdown(source)
    document = Document(docx); reference = Document(template)
    prepare_reader_styles(reference)
    report["evidence"]["sha256"] = {"source": sha256(source), "docx": sha256(docx), "pdf": sha256(pdf), "template": sha256(template)}
    report["evidence"]["paths"] = {"source": str(source.resolve()), "docx": str(docx.resolve()),
                                  "pdf": str(pdf.resolve()), "template": str(template.resolve())}
    with zipfile.ZipFile(docx) as archive:
        corrupt = archive.testzip()
        if corrupt:
            errors.append(f"Corrupt DOCX ZIP entry: {corrupt}.")
        tracked: list[str] = []
        comments: list[str] = []
        placeholders: list[str] = []
        for name in archive.namelist():
            if "comment" in name.lower() and name.lower().endswith(".xml"):
                comments.append(name)
            if not name.endswith(".xml"):
                continue
            root = ET.fromstring(archive.read(name))
            for element in root.iter():
                if element.tag in {qn("w:" + tag) for tag in ("ins", "del", "moveFrom", "moveTo", "trackRevisions")}:
                    tracked.append(name + ":" + element.tag.split("}")[-1])
                if element.tag in {qn("w:" + tag) for tag in ("commentReference", "commentRangeStart", "commentRangeEnd")}:
                    comments.append(name)
            all_text = "".join(element.text or "" for element in root.iter(qn("w:t")))
            if re.search(r"\{\{[^{}]+\}\}", all_text):
                placeholders.append(name)
        if tracked:
            errors.append(f"Revision markup/settings remain in: {sorted(set(tracked))}.")
        if comments:
            errors.append(f"Comments remain in: {sorted(set(comments))}.")
        if placeholders:
            errors.append(f"Unfilled placeholders remain in: {placeholders}.")
        report["checks"]["zip_revision_comments_slots"] = not (corrupt or tracked or comments or placeholders)
    expected = parsed.semantic_stream()
    actual = body_stream(document)
    if expected != actual:
        differences = [{"index": i + 1, "expected": a, "actual": b}
                       for i, (a, b) in enumerate(zip(expected, actual)) if a != b]
        if len(expected) != len(actual):
            errors.append(f"Body semantic block counts differ: source {len(expected)}, DOCX {len(actual)}.")
        if differences:
            errors.append(f"Body content/order differs at {len(differences)} positions; first: {differences[0]}.")
        report["evidence"]["body_differences"] = differences[:20]
    report["checks"]["source_body_order_and_text"] = expected == actual
    covers = [p for p in document.paragraphs if p.style.name == "Title"]
    if len(covers) != 1 or paragraph_xml_text(covers[0]) != parsed.title:
        errors.append("Cover title differs from the source title or is duplicated/missing.")
    contents_titles = [p for p in document.paragraphs if p.style.name == "Contents Title"]
    if len(contents_titles) != 1 or paragraph_xml_text(contents_titles[0]) != "目录":
        errors.append("Visible contents title is missing or duplicated.")
    try:
        pages = pdf_heading_pages(pdf, parsed.headings)
        report["checks"]["pdf_outline_sequence"] = True
    except ValueError as exc:
        errors.append(str(exc)); pages = []
        report["checks"]["pdf_outline_sequence"] = False
    toc = [p for p in document.paragraphs if p.style.name in CONTENTS_STYLES]
    headings = [p for p in document.paragraphs if p.style.name in HEADING_STYLES]
    all_bookmarks = list(document.element.iter(qn("w:bookmarkStart")))
    names = [x.get(qn("w:name")) for x in all_bookmarks]
    identifiers = [x.get(qn("w:id")) for x in all_bookmarks]
    if len(names) != len(set(names)) or len(identifiers) != len(set(identifiers)):
        errors.append("Duplicate bookmark names or IDs exist.")
    toc_errors: list[str] = []
    if len(toc) != len(parsed.headings) or len(headings) != len(parsed.headings):
        toc_errors.append("Source, DOCX heading and contents counts differ.")
    for i, ((level, title), entry, heading) in enumerate(zip(parsed.headings, toc, headings)):
        target = f"reader_heading_{i + 1:04d}"
        text = paragraph_xml_text(entry)
        wanted = title + "\t" + str(pages[i]) if pages else None
        if entry.style.name != f"Contents {level}" or (wanted is not None and text != wanted):
            toc_errors.append(f"Contents entry {i + 1} differs in level, title or final page: {text!r}.")
        links = entry._p.findall(qn("w:hyperlink"))
        if len(links) != 1 or links[0].get(qn("w:anchor")) != target:
            toc_errors.append(f"Contents entry {i + 1} does not jump to {target}.")
        starts = [x for x in heading._p.findall(qn("w:bookmarkStart")) if x.get(qn("w:name")) == target]
        if len(starts) != 1:
            toc_errors.append(f"Heading {i + 1} lacks its unique bookmark {target}.")
        else:
            ends = [x for x in heading._p.findall(qn("w:bookmarkEnd")) if x.get(qn("w:id")) == starts[0].get(qn("w:id"))]
            if len(ends) != 1:
                toc_errors.append(f"Heading {i + 1} has an unbalanced bookmark.")
        tabs = entry._p.findall("w:pPr/w:tabs/w:tab", NS)
        if not any(x.get(qn("w:val")) == "right" and x.get(qn("w:leader")) == "dot" for x in tabs):
            toc_errors.append(f"Contents entry {i + 1} lacks a right-aligned dotted page tab.")
    errors.extend(toc_errors)
    report["checks"]["contents_pages_levels_and_jumps"] = not toc_errors and len(pages) == len(parsed.headings)
    style_errors: list[str] = []
    font_evidence: dict[str, Any] = {}
    for name in ROLE_STYLES:
        if name not in document.styles or name not in reference.styles:
            style_errors.append(f"Missing role style: {name}."); continue
        actual_style, expected_style = document.styles[name], reference.styles[name]
        actual_fonts, expected_fonts = style_fonts(document, actual_style), style_fonts(reference, expected_style)
        font_evidence[name] = {"actual": actual_fonts, "template": expected_fonts}
        if any(value is None for value in expected_fonts.values()):
            style_errors.append(f"Template role {name} does not resolve all four font slots.")
        if actual_fonts != expected_fonts:
            style_errors.append(f"Font slots differ from the template in {name}.")
        # Exact style XML comparisons also preserve sizes, color, spacing and inheritance.
        if ET.tostring(ET.fromstring(actual_style.element.xml)) != ET.tostring(ET.fromstring(expected_style.element.xml)):
            style_errors.append(f"Style definition differs from the template: {name}.")
    font_overrides: list[str] = []
    for p_element in document.element.iter(qn("w:p")):
        paragraph = Paragraph(p_element, document._body)
        if paragraph.style.name not in ROLE_STYLES:
            continue
        expected_fonts = style_fonts(reference, reference.styles[paragraph.style.name])
        for fonts in paragraph._p.findall(".//w:rPr/w:rFonts", NS):
            if any(fonts.get(qn("w:" + slot)) for slot in THEME_SLOTS):
                font_overrides.append(paragraph.style.name + ": theme override")
            for slot in FONT_SLOTS:
                value = fonts.get(qn("w:" + slot))
                if value and value != expected_fonts[slot]:
                    font_overrides.append(f"{paragraph.style.name}: {slot}={value}")
    if font_overrides:
        style_errors.append(f"Unexpected local font overrides: {sorted(set(font_overrides))}.")
    if geometry(document) != geometry(reference):
        style_errors.append("Section/page geometry differs from the template.")
    errors.extend(style_errors)
    report["evidence"]["style_fonts"] = font_evidence
    report["checks"]["template_styles_fonts_and_geometry"] = not style_errors
    report["evidence"]["counts"] = dict(Counter(b["kind"] for b in actual))
    report["evidence"]["headings"] = len(parsed.headings)
    boundary: list[dict[str, Any]] = []
    sparse: list[dict[str, Any]] = []
    pdf_fonts: Counter[str] = Counter()
    with _open_pdf(pdf) as rendered:
        report["evidence"]["pdf_layout_backend"] = "pymupdf" if fitz is not None else "pdfplumber"
        report["evidence"]["pdf_pages"] = len(rendered)
        endings = chapter_ending_warnings(rendered, reference, parsed.headings, pages)
        report["evidence"]["sparse_chapter_endings"] = endings
        for ending in endings:
            label = "Book ending (a short closing page may be intentional)" if ending["classification"] == "book_ending" else "Chapter ending"
            warnings.append(f"{label}, page {ending['page']}: {ending['nonspace_body_characters']} nonspace body characters, "
                            f"{ending['body_lines']} rendered body lines; inspect sparse layout. Review only; do not delete text automatically.")
        for page_number, page in enumerate(rendered, 1):
            text = page.get_text()
            chars = len(re.sub(r"\s+", "", text))
            if page_number > 1 and chars < 100:
                sparse.append({"page": page_number, "nonspace_characters": chars})
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    for span in line["spans"]:
                        if not span["text"].strip():
                            continue
                        pdf_fonts[span["font"]] += len(span["text"])
                        x0, y0, x1, y1 = span["bbox"]
                        if x0 < -1 or y0 < -1 or x1 > page.rect.width + 1 or y1 > page.rect.height + 1:
                            boundary.append({"page": page_number, "text": span["text"][:80], "bbox": list(span["bbox"])})
            if "\ufffd" in text:
                warnings.append(f"Page {page_number}: extracted text contains replacement characters; inspect glyphs.")
    if boundary:
        errors.append(f"{len(boundary)} PDF text spans extend beyond the page boundary.")
    if sparse:
        warnings.append("Sparse pages require visual review; contents/chapter endings and the book ending may be valid.")
    if pdf.stat().st_mtime < docx.stat().st_mtime:
        warnings.append("PDF modification time predates the DOCX; verify that this is the final render.")
    report["evidence"]["pdf_fonts_by_text_characters"] = dict(pdf_fonts)
    report["evidence"]["sparse_pages"] = sparse
    report["evidence"]["page_boundary_anomalies"] = boundary[:50]
    report["checks"]["measurable_page_boundaries"] = not boundary
    report["ok"] = not errors
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--docx", required=True, type=Path)
    parser.add_argument("--pdf", required=True, type=Path)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    try:
        result = verify(args.source, args.docx, args.pdf, args.template)
    except (OSError, ValueError, RuntimeError, KeyError, zipfile.BadZipFile, ET.ParseError) as exc:
        result = {"ok": False, "errors": [str(exc)], "warnings": [],
                  "validation_scope": "Verification could not complete; no semantic or visual approval."}
    output = json.dumps(result, ensure_ascii=False, indent=2)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(output + "\n", encoding="utf-8")
    print(output)
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
